import io
import json
from datetime import datetime
from typing import List, Dict, Any

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_excel_bytes(records: List[Dict[str, Any]]) -> bytes:
    """
    Converts a list of extracted receipt records into a professionally styled
    Excel (.xlsx) workbook using pandas and openpyxl.
    """
    rows = []
    for idx, rec in enumerate(records, start=1):
        amount = rec.get("total_amount")
        try:
            numeric_amount = float(amount) if amount is not None and amount != "" else 0.0
        except (ValueError, TypeError):
            numeric_amount = 0.0

        rows.append(
            {
                "#": idx,
                "File Name": rec.get("filename", ""),
                "Customer / Merchant Name": rec.get("merchant_name", ""),
                "Customer (Billed To)": rec.get("customer_name", ""),
                "Invoice / Receipt #": rec.get("invoice_number", ""),
                "Date": rec.get("date", ""),
                "Currency": rec.get("currency", "USD"),
                "Total Amount": numeric_amount,
                "Extraction Method": rec.get("extraction_method", "Native PDF"),
                "Status": (rec.get("status") or "success").upper(),
            }
        )

    df = pd.DataFrame(rows)
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Extracted Receipts")
        workbook = writer.book
        worksheet = writer.sheets["Extracted Receipts"]

        # Define styles
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        body_font = Font(name="Calibri", size=11, color="1E293B")
        total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        total_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        # Style Header Row
        worksheet.row_dimensions[1].height = 28
        for col_idx in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center" if col_idx in (1, 6, 7, 10) else "left", vertical="center")
            cell.border = thin_border

        # Style Data Rows
        for row_idx in range(2, len(rows) + 2):
            worksheet.row_dimensions[row_idx].height = 22
            for col_idx in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.font = body_font
                cell.border = thin_border
                if col_idx == 8:  # Total Amount column
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_idx in (1, 6, 7, 10):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Add Summary Row if records exist
        if rows:
            summary_row_idx = len(rows) + 2
            worksheet.row_dimensions[summary_row_idx].height = 26
            worksheet.cell(row=summary_row_idx, column=2, value=f"TOTAL ({len(rows)} Receipts)").font = total_font
            total_sum = sum(r["Total Amount"] for r in rows)
            amount_cell = worksheet.cell(row=summary_row_idx, column=8, value=total_sum)
            amount_cell.font = total_font
            amount_cell.number_format = "#,##0.00"
            amount_cell.alignment = Alignment(horizontal="right", vertical="center")

            for col_idx in range(1, len(df.columns) + 1):
                c = worksheet.cell(row=summary_row_idx, column=col_idx)
                c.fill = total_fill
                c.border = thin_border

        # Auto-fit columns
        for col_idx, col_name in enumerate(df.columns, start=1):
            col_letter = get_column_letter(col_idx)
            max_len = len(str(col_name))
            for row_idx in range(2, len(rows) + 2):
                val = worksheet.cell(row=row_idx, column=col_idx).value
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            worksheet.column_dimensions[col_letter].width = min(max(max_len + 5, 12), 45)

    output.seek(0)
    return output.read()


def generate_json_string(records: List[Dict[str, Any]]) -> str:
    """
    Generates a clean, structured JSON payload with summary metadata and records.
    """
    total_amount = 0.0
    for r in records:
        try:
            if r.get("total_amount") is not None and r.get("total_amount") != "":
                total_amount += float(r["total_amount"])
        except (ValueError, TypeError):
            pass

    payload = {
        "exported_at": datetime.utcnow().isoformat() + "Z",
        "summary": {
            "total_receipts": len(records),
            "aggregate_amount": round(total_amount, 2),
        },
        "records": [
            {
                "id": r.get("id"),
                "filename": r.get("filename"),
                "merchant_name": r.get("merchant_name"),
                "customer_name": r.get("customer_name", ""),
                "invoice_number": r.get("invoice_number", ""),
                "date": r.get("date"),
                "currency": r.get("currency", "USD"),
                "total_amount": (
                    round(float(r["total_amount"]), 2)
                    if r.get("total_amount") not in (None, "")
                    else None
                ),
                "extraction_method": r.get("extraction_method", "Native PDF"),
                "confidence": r.get("confidence", 100),
                "status": r.get("status", "success"),
            }
            for r in records
        ],
    }
    return json.dumps(payload, indent=2)
